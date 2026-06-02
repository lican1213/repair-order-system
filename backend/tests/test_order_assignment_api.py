import io
import os
import tempfile
import unittest
from pathlib import Path


class OrderAssignmentApiTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmpdir = tempfile.TemporaryDirectory()
        cls.db_path = Path(cls.tmpdir.name) / "order_assignment_test.db"
        os.environ["DATABASE_URL"] = f"sqlite:///{cls.db_path}"
        os.environ["SECRET_KEY"] = "test-secret-key-for-order-assignment"
        os.environ["ADMIN_USERNAME"] = "admin"
        os.environ["ADMIN_PASSWORD"] = "ChangeThisStrongPassword123!"
        os.environ["SHOP_PHONE"] = "13800138000"
        os.environ["ORDER_WEBHOOK_ENABLED"] = "false"
        os.environ["WECOM_BOT_ENABLED"] = "false"

        from fastapi.testclient import TestClient
        from app.main import app

        cls.client = TestClient(app)
        with cls.client:
            pass

    @classmethod
    def tearDownClass(cls):
        from app.database import engine

        engine.dispose()
        cls.tmpdir.cleanup()

    def setUp(self):
        from app.auth import hash_password
        from app.database import SessionLocal
        from app.models import Order, RepairLog, UsedAppliance, User
        from app.rate_limit import limiter

        limiter._records.clear()
        db = SessionLocal()
        try:
            db.query(RepairLog).delete()
            db.query(Order).delete()
            db.query(UsedAppliance).delete()
            db.query(User).filter(User.username != "admin").delete()
            admin = db.query(User).filter(User.username == "admin").first()
            admin.role = "admin"
            admin.is_active = True
            admin.password_hash = hash_password("ChangeThisStrongPassword123!")
            db.commit()
        finally:
            db.close()

    def login_headers(self, username="admin", password="ChangeThisStrongPassword123!"):
        response = self.client.post(
            "/api/auth/login",
            json={"username": username, "password": password},
        )
        self.assertEqual(response.status_code, 200, response.text)
        token = response.json()["access_token"]
        return {"Authorization": f"Bearer {token}"}

    def create_user(self, username: str, role: str, password: str = "secret123"):
        response = self.client.post(
            "/api/auth/users",
            headers=self.login_headers(),
            json={"username": username, "password": password, "role": role},
        )
        self.assertEqual(response.status_code, 200, response.text)
        return response.json()

    def create_order(self, order_no: str, phone: str = "13100020000"):
        from app.database import SessionLocal
        from app.models import Order

        db = SessionLocal()
        try:
            order = Order(
                order_no=order_no,
                customer_name="李师傅",
                phone=phone,
                community="阳光小区",
                address="3号楼2单元501",
                service_type="维修",
                appliance_type="空调",
                fault_description="测试故障",
                status="新报修",
                followup_status="未回访",
                source="扫码报修",
            )
            db.add(order)
            db.commit()
            db.refresh(order)
            return order.id
        finally:
            db.close()

    def count_repair_logs(self) -> int:
        from app.database import SessionLocal
        from app.models import RepairLog

        db = SessionLocal()
        try:
            return db.query(RepairLog).count()
        finally:
            db.close()

    def test_admin_assigns_cancels_reassigns_and_does_not_write_repair_log(self):
        staff01 = self.create_user("staff01", "staff")
        staff02 = self.create_user("staff02", "staff")
        order_id = self.create_order("WX20260602001")
        admin_headers = self.login_headers()

        detail = self.client.get(f"/api/orders/{order_id}", headers=admin_headers)
        self.assertEqual(detail.status_code, 200, detail.text)
        self.assertIsNone(detail.json()["assigned_user_id"])
        self.assertIsNone(detail.json()["assigned_username"])

        assign_response = self.client.patch(
            f"/api/orders/{order_id}/assignment",
            headers=admin_headers,
            json={"assigned_user_id": staff01["id"]},
        )
        self.assertEqual(assign_response.status_code, 200, assign_response.text)
        self.assertEqual(assign_response.json()["assigned_user_id"], staff01["id"])
        self.assertEqual(assign_response.json()["assigned_username"], "staff01")
        self.assertEqual(self.count_repair_logs(), 0)

        duplicate_response = self.client.patch(
            f"/api/orders/{order_id}/assignment",
            headers=admin_headers,
            json={"assigned_user_id": staff01["id"]},
        )
        self.assertEqual(duplicate_response.status_code, 200, duplicate_response.text)
        self.assertEqual(duplicate_response.json()["assigned_user_id"], staff01["id"])
        self.assertEqual(self.count_repair_logs(), 0)

        reassign_response = self.client.patch(
            f"/api/orders/{order_id}/assignment",
            headers=admin_headers,
            json={"assigned_user_id": staff02["id"]},
        )
        self.assertEqual(reassign_response.status_code, 200, reassign_response.text)
        self.assertEqual(reassign_response.json()["assigned_user_id"], staff02["id"])
        self.assertEqual(reassign_response.json()["assigned_username"], "staff02")

        cancel_response = self.client.patch(
            f"/api/orders/{order_id}/assignment",
            headers=admin_headers,
            json={"assigned_user_id": None},
        )
        self.assertEqual(cancel_response.status_code, 200, cancel_response.text)
        self.assertIsNone(cancel_response.json()["assigned_user_id"])
        self.assertIsNone(cancel_response.json()["assigned_username"])
        self.assertEqual(self.count_repair_logs(), 0)

    def test_assignment_rejects_invalid_targets_and_non_admin_callers(self):
        staff01 = self.create_user("staff01", "staff")
        viewer01 = self.create_user("viewer01", "viewer")
        disabled_staff = self.create_user("disabled_staff", "staff")
        order_id = self.create_order("WX20260602002", phone="13100020002")
        admin_headers = self.login_headers()

        disable_response = self.client.patch(
            f"/api/auth/users/{disabled_staff['id']}",
            headers=admin_headers,
            json={"is_active": False},
        )
        self.assertEqual(disable_response.status_code, 200, disable_response.text)

        staff_headers = self.login_headers("staff01", "secret123")
        non_admin_response = self.client.patch(
            f"/api/orders/{order_id}/assignment",
            headers=staff_headers,
            json={"assigned_user_id": staff01["id"]},
        )
        self.assertEqual(non_admin_response.status_code, 403)

        viewer_target_response = self.client.patch(
            f"/api/orders/{order_id}/assignment",
            headers=admin_headers,
            json={"assigned_user_id": viewer01["id"]},
        )
        self.assertEqual(viewer_target_response.status_code, 400)

        disabled_target_response = self.client.patch(
            f"/api/orders/{order_id}/assignment",
            headers=admin_headers,
            json={"assigned_user_id": disabled_staff["id"]},
        )
        self.assertEqual(disabled_target_response.status_code, 400)

        missing_target_response = self.client.patch(
            f"/api/orders/{order_id}/assignment",
            headers=admin_headers,
            json={"assigned_user_id": 999999},
        )
        self.assertEqual(missing_target_response.status_code, 400)

    def test_order_visibility_is_not_filtered_by_assignee(self):
        staff01 = self.create_user("staff01", "staff")
        staff02 = self.create_user("staff02", "staff")
        own_order_id = self.create_order("WX20260602003", phone="13100020003")
        other_order_id = self.create_order("WX20260602004", phone="13100020004")
        unassigned_order_id = self.create_order("WX20260602005", phone="13100020005")
        admin_headers = self.login_headers()

        self.client.patch(
            f"/api/orders/{own_order_id}/assignment",
            headers=admin_headers,
            json={"assigned_user_id": staff01["id"]},
        )
        self.client.patch(
            f"/api/orders/{other_order_id}/assignment",
            headers=admin_headers,
            json={"assigned_user_id": staff02["id"]},
        )

        staff_headers = self.login_headers("staff01", "secret123")
        for order_id in [own_order_id, other_order_id, unassigned_order_id]:
            detail = self.client.get(f"/api/orders/{order_id}", headers=staff_headers)
            self.assertEqual(detail.status_code, 200, detail.text)

        list_response = self.client.get("/api/orders", headers=staff_headers)
        self.assertEqual(list_response.status_code, 200, list_response.text)
        visible_ids = {item["id"] for item in list_response.json()["items"]}
        self.assertEqual(visible_ids, {own_order_id, other_order_id, unassigned_order_id})

    def test_staff_write_scope_depends_on_order_assignment(self):
        staff01 = self.create_user("staff01", "staff")
        staff02 = self.create_user("staff02", "staff")
        viewer01 = self.create_user("viewer01", "viewer")
        own_order_id = self.create_order("WX20260602006", phone="13100020006")
        other_order_id = self.create_order("WX20260602007", phone="13100020007")
        unassigned_order_id = self.create_order("WX20260602008", phone="13100020008")
        admin_headers = self.login_headers()

        self.client.patch(
            f"/api/orders/{own_order_id}/assignment",
            headers=admin_headers,
            json={"assigned_user_id": staff01["id"]},
        )
        self.client.patch(
            f"/api/orders/{other_order_id}/assignment",
            headers=admin_headers,
            json={"assigned_user_id": staff02["id"]},
        )

        staff_headers = self.login_headers("staff01", "secret123")
        own_patch = self.client.patch(
            f"/api/orders/{own_order_id}",
            headers=staff_headers,
            json={"status": "已联系"},
        )
        self.assertEqual(own_patch.status_code, 200, own_patch.text)

        unassigned_patch = self.client.patch(
            f"/api/orders/{unassigned_order_id}",
            headers=staff_headers,
            json={"status": "已联系"},
        )
        self.assertEqual(unassigned_patch.status_code, 200, unassigned_patch.text)
        self.assertIsNone(unassigned_patch.json()["assigned_user_id"])

        other_patch = self.client.patch(
            f"/api/orders/{other_order_id}",
            headers=staff_headers,
            json={"status": "已联系"},
        )
        self.assertEqual(other_patch.status_code, 403)

        viewer_headers = self.login_headers("viewer01", "secret123")
        viewer_patch = self.client.patch(
            f"/api/orders/{own_order_id}",
            headers=viewer_headers,
            json={"status": "已预约"},
        )
        self.assertEqual(viewer_patch.status_code, 403)

        admin_patch = self.client.patch(
            f"/api/orders/{other_order_id}",
            headers=admin_headers,
            json={"status": "已联系"},
        )
        self.assertEqual(admin_patch.status_code, 200, admin_patch.text)

    def test_assignee_filters_export_public_schema_and_delete_user_cleanup(self):
        staff01 = self.create_user("staff01", "staff")
        staff02 = self.create_user("staff02", "staff")
        assigned_order_id = self.create_order("WX20260602009", phone="13100020009")
        other_order_id = self.create_order("WX20260602010", phone="13100020010")
        unassigned_order_id = self.create_order("WX20260602011", phone="13100020011")
        admin_headers = self.login_headers()

        self.client.patch(
            f"/api/orders/{assigned_order_id}/assignment",
            headers=admin_headers,
            json={"assigned_user_id": staff01["id"]},
        )
        self.client.patch(
            f"/api/orders/{other_order_id}/assignment",
            headers=admin_headers,
            json={"assigned_user_id": staff02["id"]},
        )

        staff_filter = self.client.get(
            f"/api/orders?assignee={staff01['id']}",
            headers=admin_headers,
        )
        self.assertEqual(staff_filter.status_code, 200, staff_filter.text)
        self.assertEqual(
            {item["id"] for item in staff_filter.json()["items"]},
            {assigned_order_id},
        )

        unassigned_filter = self.client.get(
            "/api/orders?assignee=unassigned",
            headers=admin_headers,
        )
        self.assertEqual(unassigned_filter.status_code, 200, unassigned_filter.text)
        self.assertEqual(
            {item["id"] for item in unassigned_filter.json()["items"]},
            {unassigned_order_id},
        )

        mine_filter = self.client.get(
            "/api/orders?assignee=mine",
            headers=self.login_headers("staff01", "secret123"),
        )
        self.assertEqual(mine_filter.status_code, 200, mine_filter.text)
        self.assertEqual(
            {item["id"] for item in mine_filter.json()["items"]},
            {assigned_order_id},
        )

        export_response = self.client.get("/api/export/orders", headers=admin_headers)
        self.assertEqual(export_response.status_code, 200, export_response.text)
        from openpyxl import load_workbook

        workbook = load_workbook(io.BytesIO(export_response.content), read_only=True)
        headers = [cell.value for cell in next(workbook.active.iter_rows(max_row=1))]
        self.assertIn("负责人", headers)

        public_response = self.client.post(
            "/api/public/submit",
            headers={"x-forwarded-for": "203.0.113.25"},
            json={
                "customer_name": "王女士",
                "phone": "13100020012",
                "community": "幸福花园",
                "address": "1号楼101",
                "service_type": "清洗",
                "appliance_type": "油烟机",
                "fault_description": "需要清洗",
            },
        )
        self.assertEqual(public_response.status_code, 200, public_response.text)
        self.assertNotIn("assigned_user_id", public_response.json())
        self.assertNotIn("assigned_username", public_response.json())

        notification_response = self.client.get(
            "/api/orders/notifications/new?after_id=0",
            headers=admin_headers,
        )
        self.assertEqual(notification_response.status_code, 200, notification_response.text)
        for item in notification_response.json()["orders"]:
            self.assertNotIn("assigned_user_id", item)
            self.assertNotIn("assigned_username", item)

        delete_response = self.client.delete(
            f"/api/auth/users/{staff01['id']}",
            headers=admin_headers,
        )
        self.assertEqual(delete_response.status_code, 200, delete_response.text)
        cleaned_detail = self.client.get(
            f"/api/orders/{assigned_order_id}",
            headers=admin_headers,
        )
        self.assertEqual(cleaned_detail.status_code, 200, cleaned_detail.text)
        self.assertIsNone(cleaned_detail.json()["assigned_user_id"])
        self.assertIsNone(cleaned_detail.json()["assigned_username"])


if __name__ == "__main__":
    unittest.main()
